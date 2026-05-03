from flask import Flask, request, jsonify, make_response
from flask_sqlalchemy import SQLAlchemy
from flask_cors import CORS
from datetime import datetime
import json
import os
import sys
import traceback
from io import BytesIO

HAS_OPENPYXL = False
OPENPYXL_ERROR = None
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    HAS_OPENPYXL = True
except ImportError as e:
    HAS_OPENPYXL = False
    OPENPYXL_ERROR = str(e)
    print(f"openpyxl导入失败: {e}", file=sys.stderr)
    print(f"Python路径: {sys.executable}", file=sys.stderr)

app = Flask(__name__)
basedir = os.path.abspath(os.path.dirname(__file__))
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///' + os.path.join(basedir, 'test_platform.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db = SQLAlchemy(app)
CORS(app)


class TestCaseDirectory(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    parent_id = db.Column(db.Integer, db.ForeignKey('test_case_directory.id'), nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.now)
    updated_at = db.Column(db.DateTime, default=datetime.now, onupdate=datetime.now)
    
    parent = db.relationship('TestCaseDirectory', remote_side=[id], backref='children')

    def to_dict(self, include_children=True):
        result = {
            'id': self.id,
            'name': self.name,
            'parent_id': self.parent_id,
            'created_at': self.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'updated_at': self.updated_at.strftime('%Y-%m-%d %H:%M:%S')
        }
        if include_children:
            result['children'] = [child.to_dict(include_children=True) for child in self.children]
        return result


class TestCase(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    preconditions = db.Column(db.Text, nullable=True)
    steps = db.Column(db.Text, nullable=False)
    expected_results = db.Column(db.Text, nullable=False)
    description = db.Column(db.Text, nullable=True)
    priority = db.Column(db.String(2), default='P0')
    directory_id = db.Column(db.Integer, db.ForeignKey('test_case_directory.id'), nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.now)
    updated_at = db.Column(db.DateTime, default=datetime.now, onupdate=datetime.now)
    
    directory = db.relationship('TestCaseDirectory', backref='test_cases')

    def to_dict(self):
        try:
            priority = self.priority
        except AttributeError:
            priority = 'P0'
        
        return {
            'id': self.id,
            'name': self.name,
            'preconditions': self.preconditions,
            'steps': json.loads(self.steps) if self.steps else [],
            'expected_results': json.loads(self.expected_results) if self.expected_results else [],
            'description': getattr(self, 'description', None),
            'priority': priority,
            'directory_id': self.directory_id,
            'directory_name': self.directory.name if self.directory else None,
            'created_at': self.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'updated_at': self.updated_at.strftime('%Y-%m-%d %H:%M:%S')
        }


class TestPlan(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.now)
    updated_at = db.Column(db.DateTime, default=datetime.now, onupdate=datetime.now)

    def to_dict(self):
        test_cases = TestPlanTestCase.query.filter_by(test_plan_id=self.id).order_by(
            TestPlanTestCase.order
        ).all()
        test_case_ids = [tc.test_case_id for tc in test_cases]
        test_cases_data = []
        latest_case_updated_at = None
        for tc_id in test_case_ids:
            tc = TestCase.query.get(tc_id)
            if tc:
                test_cases_data.append(tc.to_dict())
                if latest_case_updated_at is None or tc.updated_at > latest_case_updated_at:
                    latest_case_updated_at = tc.updated_at

        latest_updated_at = self.updated_at
        if latest_case_updated_at is not None and latest_case_updated_at > latest_updated_at:
            latest_updated_at = latest_case_updated_at

        if latest_updated_at < self.created_at:
            latest_updated_at = self.created_at

        return {
            'id': self.id,
            'name': self.name,
            'test_cases': test_cases_data,
            'created_at': self.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'updated_at': latest_updated_at.strftime('%Y-%m-%d %H:%M:%S')
        }


class TestPlanTestCase(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    test_plan_id = db.Column(db.Integer, db.ForeignKey('test_plan.id'), nullable=False)
    test_case_id = db.Column(db.Integer, db.ForeignKey('test_case.id'), nullable=False)
    order = db.Column(db.Integer, default=0)
    created_at = db.Column(db.DateTime, default=datetime.now)


@app.route('/api/directories', methods=['GET'])
def get_directories():
    directories = TestCaseDirectory.query.filter_by(parent_id=None).all()
    return jsonify([d.to_dict(include_children=True) for d in directories])


@app.route('/api/directories/all', methods=['GET'])
def get_all_directories():
    directories = TestCaseDirectory.query.order_by(TestCaseDirectory.created_at.asc()).all()
    return jsonify([{
        'id': d.id,
        'name': d.name,
        'parent_id': d.parent_id
    } for d in directories])


@app.route('/api/directories', methods=['POST'])
def create_directory():
    data = request.get_json()
    
    if not data.get('name') or len(data.get('name', '').strip()) == 0:
        return jsonify({'error': '目录名称不能为空'}), 400
    
    if len(data.get('name', '')) > 100:
        return jsonify({'error': '目录名称不能超过100字'}), 400
    
    parent_id = data.get('parent_id')
    if parent_id is not None:
        parent = TestCaseDirectory.query.get(parent_id)
        if not parent:
            return jsonify({'error': '父目录不存在'}), 400
    
    directory = TestCaseDirectory(
        name=data['name'].strip(),
        parent_id=parent_id
    )
    db.session.add(directory)
    db.session.commit()
    
    return jsonify(directory.to_dict(include_children=False)), 201


@app.route('/api/directories/<int:id>', methods=['PUT'])
def update_directory(id):
    directory = TestCaseDirectory.query.get_or_404(id)
    data = request.get_json()
    
    if not data.get('name') or len(data.get('name', '').strip()) == 0:
        return jsonify({'error': '目录名称不能为空'}), 400
    
    if len(data.get('name', '')) > 100:
        return jsonify({'error': '目录名称不能超过100字'}), 400
    
    parent_id = data.get('parent_id')
    if parent_id is not None:
        if parent_id == id:
            return jsonify({'error': '不能将目录设置为自己的子目录'}), 400
        parent = TestCaseDirectory.query.get(parent_id)
        if not parent:
            return jsonify({'error': '父目录不存在'}), 400
    
    directory.name = data['name'].strip()
    if 'parent_id' in data:
        directory.parent_id = parent_id
    
    db.session.commit()
    
    return jsonify(directory.to_dict(include_children=False))


@app.route('/api/directories/<int:id>', methods=['DELETE'])
def delete_directory(id):
    directory = TestCaseDirectory.query.get_or_404(id)
    
    test_cases = TestCase.query.filter_by(directory_id=id).all()
    if test_cases:
        return jsonify({'error': '目录下存在测试用例，无法删除'}), 400
    
    if directory.children:
        return jsonify({'error': '目录下存在子目录，无法删除'}), 400
    
    db.session.delete(directory)
    db.session.commit()
    
    return jsonify({'message': '删除成功'})


def get_all_subdirectory_ids(directory_id):
    """获取指定目录及其所有子目录的ID列表"""
    if directory_id is None:
        return None
    
    result = [directory_id]
    
    def get_children(parent_id):
        children = TestCaseDirectory.query.filter_by(parent_id=parent_id).all()
        for child in children:
            result.append(child.id)
            get_children(child.id)
    
    get_children(directory_id)
    return result


def get_directory_depth(directory_id):
    """获取目录的深度（根目录深度为1）"""
    if directory_id is None:
        return 0
    
    depth = 1
    directory = TestCaseDirectory.query.get(directory_id)
    if not directory:
        return 0
    
    while directory.parent_id is not None:
        depth += 1
        directory = TestCaseDirectory.query.get(directory.parent_id)
        if not directory:
            break
    
    return depth


def is_descendant(ancestor_id, descendant_id):
    """检查一个目录是否是另一个目录的后代"""
    if ancestor_id is None or descendant_id is None:
        return False
    if ancestor_id == descendant_id:
        return False
    
    descendant = TestCaseDirectory.query.get(descendant_id)
    if not descendant:
        return False
    
    while descendant.parent_id is not None:
        if descendant.parent_id == ancestor_id:
            return True
        descendant = TestCaseDirectory.query.get(descendant.parent_id)
        if not descendant:
            break
    
    return False


def get_directory_path(directory_id):
    """获取目录的完整路径（如：根目录/子目录/孙子目录）"""
    if directory_id is None:
        return ''
    
    path = []
    directory = TestCaseDirectory.query.get(directory_id)
    if not directory:
        return ''
    
    while directory:
        path.insert(0, directory.name)
        if directory.parent_id is None:
            break
        directory = TestCaseDirectory.query.get(directory.parent_id)
    
    return '/'.join(path)


@app.route('/api/directories/<int:id>/move', methods=['POST'])
def move_directory(id):
    """移动目录到新的父目录"""
    directory = TestCaseDirectory.query.get_or_404(id)
    data = request.get_json()
    
    new_parent_id = data.get('new_parent_id')
    
    if new_parent_id is not None:
        if new_parent_id == id:
            return jsonify({'error': '不能将目录移动到自己下面'}), 400
        
        new_parent = TestCaseDirectory.query.get(new_parent_id)
        if not new_parent:
            return jsonify({'error': '目标目录不存在'}), 400
        
        if is_descendant(id, new_parent_id):
            return jsonify({'error': '不能将目录移动到自己的子目录下面'}), 400
    
    directory.parent_id = new_parent_id
    db.session.commit()
    
    return jsonify({
        'message': '目录移动成功',
        'directory': directory.to_dict(include_children=False)
    })


@app.route('/api/testcases/batch/move', methods=['POST'])
def batch_move_testcases():
    """批量移动测试用例到目标目录"""
    data = request.get_json()
    
    test_case_ids = data.get('test_case_ids', [])
    target_directory_id = data.get('target_directory_id')
    
    if not test_case_ids or len(test_case_ids) == 0:
        return jsonify({'error': '请选择至少一个测试用例'}), 400
    
    if target_directory_id is not None:
        target_directory = TestCaseDirectory.query.get(target_directory_id)
        if not target_directory:
            return jsonify({'error': '目标目录不存在'}), 400
    
    test_cases = TestCase.query.filter(TestCase.id.in_(test_case_ids)).all()
    for test_case in test_cases:
        test_case.directory_id = target_directory_id
    
    db.session.commit()
    
    return jsonify({
        'message': f'成功移动 {len(test_cases)} 个测试用例',
        'moved_count': len(test_cases)
    })


def get_column_letter(n):
    """将数字转换为Excel列字母 (1=A, 2=B, ..., 26=Z, 27=AA, etc.)"""
    result = []
    while n > 0:
        n -= 1
        result.append(chr(ord('A') + (n % 26)))
        n = n // 26
    return ''.join(reversed(result))


def safe_str(value):
    """安全转换为字符串，处理None和特殊字符"""
    if value is None:
        return ''
    try:
        return str(value)
    except Exception:
        return ''


def create_excel_report(test_cases):
    """创建Excel报告"""
    if not HAS_OPENPYXL:
        print("openpyxl未安装，无法创建Excel报告", file=sys.stderr)
        return None
    
    try:
        print(f"开始创建Excel报告，共 {len(test_cases)} 条测试用例", file=sys.stderr)
        
        wb = Workbook()
        ws = wb.active
        ws.title = '测试用例'
        
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font_white = Font(bold=True, size=12, color='FFFFFF')
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['用例ID', '用例名称', '用例等级', '前置条件', '测试步骤', '预期结果', '描述', '目录', '创建时间', '编辑时间']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = center_alignment
            cell.border = thin_border
        
        for row_num, test_case in enumerate(test_cases, 2):
            try:
                steps = test_case.get('steps', [])
                expected_results = test_case.get('expected_results', [])
                
                if not isinstance(steps, list):
                    steps = []
                if not isinstance(expected_results, list):
                    expected_results = []
                
                steps_text = ''
                for i, step in enumerate(steps, 1):
                    step_str = safe_str(step)
                    steps_text += f'{i}. {step_str}\n'
                
                expected_text = ''
                for i, result in enumerate(expected_results, 1):
                    result_str = safe_str(result)
                    expected_text += f'{i}. {result_str}\n'
                
                row_data = [
                    test_case.get('id', ''),
                    safe_str(test_case.get('name', '')),
                    safe_str(test_case.get('priority', 'P0')),
                    safe_str(test_case.get('preconditions', '')),
                    steps_text.strip(),
                    expected_text.strip(),
                    safe_str(test_case.get('description', '')),
                    safe_str(test_case.get('directory_name', '未分配')),
                    safe_str(test_case.get('created_at', '')),
                    safe_str(test_case.get('updated_at', ''))
                ]
                
                for col_num, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_num, column=col_num, value=value)
                    cell.alignment = left_alignment if col_num in [2, 4, 5, 6, 7, 8] else center_alignment
                    cell.border = thin_border
            except Exception as row_error:
                print(f"处理第 {row_num} 行数据时出错: {str(row_error)}", file=sys.stderr)
                traceback.print_exc(file=sys.stderr)
                continue
        
        column_widths = [8, 30, 10, 20, 40, 40, 20, 20, 20, 20]
        for i, width in enumerate(column_widths, 1):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = width
        
        print("Excel报告创建成功", file=sys.stderr)
        return wb
    except Exception as e:
        print(f"创建Excel报告失败: {str(e)}", file=sys.stderr)
        traceback.print_exc(file=sys.stderr)
        return None


@app.route('/api/testcases/export', methods=['POST'])
def export_testcases():
    """导出测试用例为Excel"""
    print("收到导出请求", file=sys.stderr)
    
    if not HAS_OPENPYXL:
        error_msg = '服务器未安装Excel导出依赖（openpyxl）。'
        if OPENPYXL_ERROR:
            error_msg += f' 错误详情: {OPENPYXL_ERROR}'
        print(f"导出错误: {error_msg}", file=sys.stderr)
        return jsonify({'error': error_msg}), 500
    
    try:
        data = request.get_json(force=True, silent=True)
        if data is None:
            print("无法解析请求JSON", file=sys.stderr)
            return jsonify({'error': '请求格式错误'}), 400
        
        test_case_ids = data.get('test_case_ids', [])
        print(f"导出请求包含 {len(test_case_ids)} 个测试用例ID", file=sys.stderr)
        
        if not test_case_ids or len(test_case_ids) == 0:
            return jsonify({'error': '请选择至少一个测试用例'}), 400
        
        test_cases = TestCase.query.filter(TestCase.id.in_(test_case_ids)).order_by(TestCase.created_at.desc()).all()
        print(f"从数据库查询到 {len(test_cases)} 条测试用例", file=sys.stderr)
        
        test_cases_data = []
        for tc in test_cases:
            try:
                tc_dict = tc.to_dict()
                test_cases_data.append(tc_dict)
            except Exception as dict_error:
                print(f"转换测试用例 {tc.id} 为字典时出错: {str(dict_error)}", file=sys.stderr)
                traceback.print_exc(file=sys.stderr)
        
        print(f"准备创建Excel报告，共 {len(test_cases_data)} 条有效数据", file=sys.stderr)
        
        wb = create_excel_report(test_cases_data)
        if not wb:
            print("创建Excel报告失败，返回错误响应", file=sys.stderr)
            return jsonify({'error': '创建Excel报告失败'}), 500
        
        print("Excel报告创建成功，准备保存到内存", file=sys.stderr)
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        file_content = output.getvalue()
        file_size = len(file_content)
        print(f"Excel文件大小: {file_size} 字节", file=sys.stderr)
        
        response = make_response(file_content)
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Length'] = str(file_size)
        response.headers['Content-Disposition'] = f'attachment; filename=测试用例_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
        
        print("导出响应准备完成，返回给客户端", file=sys.stderr)
        return response
    except Exception as e:
        print(f"导出Excel失败: {str(e)}", file=sys.stderr)
        traceback.print_exc(file=sys.stderr)
        try:
            return jsonify({'error': f'导出失败: {str(e)}'}), 500
        except Exception as json_error:
            print(f"构建错误响应时也失败: {str(json_error)}", file=sys.stderr)
            return make_response(
                f'{{"error": "导出失败: {str(e)}"}}', 
                500, 
                {'Content-Type': 'application/json'}
            )


@app.route('/api/testcases/all', methods=['GET'])
def get_all_testcases():
    directory_id = request.args.get('directory_id', type=int)
    include_subdirs = request.args.get('include_subdirs', 'true', type=str).lower() == 'true'
    
    query = TestCase.query
    if directory_id is not None:
        if include_subdirs:
            directory_ids = get_all_subdirectory_ids(directory_id)
            query = query.filter(TestCase.directory_id.in_(directory_ids))
        else:
            query = query.filter_by(directory_id=directory_id)
    
    test_cases = query.order_by(TestCase.created_at.desc()).all()
    return jsonify([tc.to_dict() for tc in test_cases])


@app.route('/api/testcases', methods=['GET'])
def get_testcases():
    page = request.args.get('page', 1, type=int)
    page_size = request.args.get('page_size', 10, type=int)
    directory_id = request.args.get('directory_id', type=int)
    include_subdirs = request.args.get('include_subdirs', 'true', type=str).lower() == 'true'
    
    valid_page_sizes = [10, 20, 50, 100]
    if page_size not in valid_page_sizes:
        page_size = 10
    
    query = TestCase.query
    if directory_id is not None:
        if include_subdirs:
            directory_ids = get_all_subdirectory_ids(directory_id)
            query = query.filter(TestCase.directory_id.in_(directory_ids))
        else:
            query = query.filter_by(directory_id=directory_id)
    
    pagination = query.order_by(TestCase.created_at.desc()).paginate(
        page=page, per_page=page_size, error_out=False
    )
    
    return jsonify({
        'testcases': [tc.to_dict() for tc in pagination.items],
        'total': pagination.total,
        'page': page,
        'page_size': page_size,
        'total_pages': pagination.pages
    })


@app.route('/api/testcases', methods=['POST'])
def create_testcase():
    data = request.get_json()
    
    if not data.get('name') or len(data.get('name', '').strip()) == 0:
        return jsonify({'error': '用例名称不能为空'}), 400
    
    if len(data.get('name', '')) > 200:
        return jsonify({'error': '用例名称不能超过200字'}), 400
    
    directory_id = data.get('directory_id')
    if directory_id is None:
        return jsonify({'error': '请选择目录'}), 400
    
    if directory_id is not None:
        directory = TestCaseDirectory.query.get(directory_id)
        if not directory:
            return jsonify({'error': '目录不存在'}), 400
    
    steps = data.get('steps', [])
    expected_results = data.get('expected_results', [])
    
    if len(steps) != len(expected_results):
        return jsonify({'error': '测试步骤数量与预期结果数量不一致'}), 400
    
    for i, step in enumerate(steps):
        if len(step) > 500:
            return jsonify({'error': f'第{i+1}步测试步骤不能超过500字'}), 400
    
    for i, result in enumerate(expected_results):
        if len(result) > 500:
            return jsonify({'error': f'第{i+1}步预期结果不能超过500字'}), 400
    
    testcase_kwargs = {
        'name': data['name'],
        'preconditions': data.get('preconditions', ''),
        'steps': json.dumps(steps),
        'expected_results': json.dumps(expected_results),
        'directory_id': directory_id
    }
    
    # 处理 description 字段
    try:
        testcase_test = TestCase(**testcase_kwargs)
        _ = testcase_test.description
        # 如果可以访问 description 字段
        description = data.get('description', '')
        if len(description) > 200:
            return jsonify({'error': '描述不能超过200字'}), 400
        testcase_kwargs['description'] = description
    except AttributeError:
        pass
    
    # 处理 priority 字段
    try:
        testcase_test = TestCase(**testcase_kwargs)
        _ = testcase_test.priority
        # 如果可以访问 priority 字段
        priority = data.get('priority', 'P0')
        if priority not in ['P0', 'P1', 'P2', 'P3', 'P4']:
            priority = 'P0'
        testcase_kwargs['priority'] = priority
    except AttributeError:
        pass
    
    testcase = TestCase(**testcase_kwargs)
    
    db.session.add(testcase)
    db.session.commit()
    
    return jsonify(testcase.to_dict()), 201


@app.route('/api/testcases/<int:id>', methods=['GET'])
def get_testcase(id):
    testcase = TestCase.query.get_or_404(id)
    return jsonify(testcase.to_dict())


@app.route('/api/testcases/<int:id>', methods=['PUT'])
def update_testcase(id):
    testcase = TestCase.query.get_or_404(id)
    data = request.get_json()
    
    if not data.get('name') or len(data.get('name', '').strip()) == 0:
        return jsonify({'error': '用例名称不能为空'}), 400
    
    if len(data.get('name', '')) > 200:
        return jsonify({'error': '用例名称不能超过200字'}), 400
    
    directory_id = data.get('directory_id')
    if directory_id is None:
        return jsonify({'error': '请选择目录'}), 400
    
    if directory_id is not None:
        directory = TestCaseDirectory.query.get(directory_id)
        if not directory:
            return jsonify({'error': '目录不存在'}), 400
    
    steps = data.get('steps', [])
    expected_results = data.get('expected_results', [])
    
    if len(steps) != len(expected_results):
        return jsonify({'error': '测试步骤数量与预期结果数量不一致'}), 400
    
    for i, step in enumerate(steps):
        if len(step) > 500:
            return jsonify({'error': f'第{i+1}步测试步骤不能超过500字'}), 400
    
    for i, result in enumerate(expected_results):
        if len(result) > 500:
            return jsonify({'error': f'第{i+1}步预期结果不能超过500字'}), 400
    
    testcase.name = data['name']
    testcase.preconditions = data.get('preconditions', '')
    testcase.steps = json.dumps(steps)
    testcase.expected_results = json.dumps(expected_results)
    testcase.directory_id = directory_id
    
    # 处理 description 字段
    try:
        _ = testcase.description
        description = data.get('description', '')
        if len(description) > 200:
            return jsonify({'error': '描述不能超过200字'}), 400
        testcase.description = description
    except AttributeError:
        pass
    
    # 处理 priority 字段
    try:
        _ = testcase.priority
        priority = data.get('priority', 'P0')
        if priority not in ['P0', 'P1', 'P2', 'P3', 'P4']:
            priority = 'P0'
        testcase.priority = priority
    except AttributeError:
        pass
    
    db.session.commit()
    
    return jsonify(testcase.to_dict())


@app.route('/api/testplans', methods=['GET'])
def get_testplans():
    page = request.args.get('page', 1, type=int)
    page_size = request.args.get('page_size', 10, type=int)
    
    valid_page_sizes = [10, 20, 50, 100]
    if page_size not in valid_page_sizes:
        page_size = 10
    
    pagination = TestPlan.query.order_by(TestPlan.created_at.desc()).paginate(
        page=page, per_page=page_size, error_out=False
    )
    
    return jsonify({
        'testplans': [tp.to_dict() for tp in pagination.items],
        'total': pagination.total,
        'page': page,
        'page_size': page_size,
        'total_pages': pagination.pages
    })


@app.route('/api/testplans', methods=['POST'])
def create_testplan():
    data = request.get_json()
    
    if not data.get('name') or len(data.get('name', '').strip()) == 0:
        return jsonify({'error': '测试计划名称不能为空'}), 400
    
    if len(data.get('name', '')) > 200:
        return jsonify({'error': '测试计划名称不能超过200字'}), 400
    
    test_case_ids = data.get('test_case_ids', [])
    if not test_case_ids or len(test_case_ids) == 0:
        return jsonify({'error': '请至少选择一个测试用例'}), 400
    
    testplan = TestPlan(
        name=data['name'].strip()
    )
    db.session.add(testplan)
    db.session.flush()
    
    for order, tc_id in enumerate(test_case_ids):
        relation = TestPlanTestCase(
            test_plan_id=testplan.id,
            test_case_id=tc_id,
            order=order
        )
        db.session.add(relation)
    
    db.session.commit()
    
    return jsonify(testplan.to_dict()), 201


@app.route('/api/testplans/<int:id>', methods=['GET'])
def get_testplan(id):
    testplan = TestPlan.query.get_or_404(id)
    return jsonify(testplan.to_dict())


@app.route('/api/testplans/<int:id>', methods=['PUT'])
def update_testplan(id):
    testplan = TestPlan.query.get_or_404(id)
    data = request.get_json()
    
    if not data.get('name') or len(data.get('name', '').strip()) == 0:
        return jsonify({'error': '测试计划名称不能为空'}), 400
    
    if len(data.get('name', '')) > 200:
        return jsonify({'error': '测试计划名称不能超过200字'}), 400
    
    test_case_ids = data.get('test_case_ids', [])
    
    testplan.name = data['name'].strip()
    
    TestPlanTestCase.query.filter_by(test_plan_id=testplan.id).delete()
    
    for order, tc_id in enumerate(test_case_ids):
        relation = TestPlanTestCase(
            test_plan_id=testplan.id,
            test_case_id=tc_id,
            order=order
        )
        db.session.add(relation)
    
    db.session.commit()
    
    return jsonify(testplan.to_dict())


if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(debug=True, port=6000)
